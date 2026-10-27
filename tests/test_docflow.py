import importlib.util
import json
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
from unittest.mock import patch

SCRIPT = Path(__file__).resolve().parents[1] / 'skills/organize-search-documents/scripts/docflow.py'
spec = importlib.util.spec_from_file_location('docflow', SCRIPT)
d = importlib.util.module_from_spec(spec)
spec.loader.exec_module(d)


class WorkflowTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.base = Path(self.tmp.name).resolve()
        self.root = self.base / 'input'
        self.root.mkdir()
        self.db = self.base / 'index.sqlite3'
        self.file = self.root / '合同.txt'
        self.file.write_text('模拟资料\n付款期限为30天。\n', encoding='utf-8')

    def tearDown(self):
        self.tmp.cleanup()

    def sync(self):
        return d.sync(self.root, self.db)

    def body(self):
        return d.snapshot(self.db)

    def test_chinese_locations_and_no_hit(self):
        self.sync()
        found = d.search(self.body(), '付款 30天')['results']
        self.assertEqual(len(found), 1)
        self.assertEqual(found[0]['location'], 'lines 1-2')
        self.assertEqual(d.search(self.body(), '不存在')['total'], 0)

    def test_dedup_and_reuse(self):
        (self.root / '副本.txt').write_bytes(self.file.read_bytes())
        self.assertEqual(len(self.sync()['duplicates'][0]), 2)
        self.sync()
        self.assertTrue(all(f.get('reused') for f in self.body()['files']))

    def test_changed_same_size_and_removed(self):
        self.sync()
        self.file.write_text('模拟资料\n付款期限为60天。\n', encoding='utf-8')
        self.sync()
        self.assertEqual(d.search(self.body(), '30天')['total'], 0)
        self.assertEqual(d.search(self.body(), '60天')['total'], 1)
        self.file.unlink()
        self.sync()
        self.assertEqual(d.search(self.body(), '付款')['total'], 0)

    def test_errors_unsupported_and_exit_code(self):
        (self.root / 'bad.pdf').write_bytes(b'not a pdf')
        (self.root / 'old.xls').write_bytes(b'simulated unsupported, not real XLS')
        p = subprocess.run([sys.executable, str(SCRIPT), 'sync', '--root', str(self.root), '--db', str(self.db)], capture_output=True, text=True)
        self.assertEqual(p.returncode, 2)
        self.assertEqual(json.loads(p.stdout)['counts'], {'ok': 1, 'partial': 0, 'error': 1, 'unsupported': 1})

    def test_hidden_and_symlink(self):
        (self.root / '.secret.txt').write_text('secret')
        try:
            (self.root / 'link.txt').symlink_to(self.file)
        except OSError:
            self.skipTest('symlink not available')
        self.sync()
        self.assertEqual(len(self.body()['files']), 1)
        self.assertEqual(len(self.body()['skipped']), 2)

    def test_cross_chunk_and_pagination(self):
        self.file.write_text('A' * 1998 + '付款期限' + 'B' * 2500, encoding='utf-8')
        self.sync()
        self.assertGreater(d.search(self.body(), '付款期限')['total'], 0)
        result = d.search(self.body(), 'B', limit=1)
        self.assertTrue(result['has_more'])
        self.assertEqual(len(d.search(self.body(), 'B', limit=1, offset=1)['results']), 1)

    def test_export_no_overwrite_no_absolute_root(self):
        self.sync()
        out = self.base / 'bundle'
        d.export_bundle(self.body(), out, '付款')
        text = (out / 'evidence.txt').read_text(encoding='utf-8')
        self.assertIn('合同.txt | lines', text)
        self.assertNotIn(str(self.root), (out / 'evidence.json').read_text(encoding='utf-8'))
        with self.assertRaises(FileExistsError):
            d.export_bundle(self.body(), out, '付款')

    def test_archive_preview_copy_repeat_and_conflict(self):
        self.sync()
        original = self.file.read_bytes()
        target = self.base / 'archive'
        plan = d.archive(self.body(), target)
        self.assertFalse(target.exists())
        self.assertEqual(plan['plan'][0]['action'], 'would_copy')
        copied = d.archive(self.body(), target, True)['plan'][0]
        self.assertEqual(Path(copied['target']).read_bytes(), original)
        self.assertEqual(d.archive(self.body(), target, True)['plan'][0]['action'], 'skip_identical')
        Path(copied['target']).write_text('do not overwrite')
        self.assertEqual(d.archive(self.body(), target, True)['plan'][0]['action'], 'conflict')
        self.assertEqual(Path(copied['target']).read_text(), 'do not overwrite')
        self.assertEqual(self.file.read_bytes(), original)

    def test_archive_stale_and_overlap(self):
        self.sync()
        with self.assertRaises(ValueError): d.archive(self.body(), self.root / 'archive')
        self.file.write_text('changed')
        result = d.archive(self.body(), self.base / 'archive', True)
        self.assertEqual(result['plan'][0]['reason'], 'source_changed_run_sync')

    def test_database_scope_and_failed_scan(self):
        self.sync()
        before = self.db.read_bytes()
        with patch.object(d, 'scan', side_effect=PermissionError('incomplete enumeration')):
            with self.assertRaises(PermissionError): self.sync()
        self.assertEqual(self.db.read_bytes(), before)
        with self.assertRaises(ValueError): d.sync(self.root, self.root / 'bad.sqlite3')
        other = self.base / 'other'
        other.mkdir()
        with self.assertRaises(ValueError): d.sync(other, self.db)

    def test_separate_process_persistence(self):
        self.sync()
        p = subprocess.run([sys.executable, str(SCRIPT), 'search', '--db', str(self.db), '--query', '付款'], capture_output=True, text=True)
        self.assertEqual(p.returncode, 0)
        self.assertEqual(json.loads(p.stdout)['total'], 1)

    def test_csv_multiline_and_legacy_encoding(self):
        (self.root / 'table.csv').write_text('名称,说明\n模拟,"第一行\n第二行"\n', encoding='utf-8')
        (self.root / 'legacy.txt').write_bytes('中文编码'.encode('gb18030'))
        self.sync()
        self.assertEqual(d.search(self.body(), '第二行')['results'][0]['location'], 'CSV record 2')
        self.assertIn('GB18030', d.search(self.body(), '中文')['results'][0]['warnings'][0])

    def test_docx_paragraph_table(self):
        from docx import Document
        doc = Document()
        doc.add_paragraph('模拟审批流程')
        doc.add_table(rows=1, cols=1).cell(0, 0).text = '模拟付款节点'
        doc.save(self.root / 'sample.docx')
        self.sync()
        self.assertEqual(d.search(self.body(), '审批')['results'][0]['location'], 'paragraph 1')
        self.assertEqual(d.search(self.body(), '节点')['results'][0]['location'], 'table 1 row 1 column 1')

    def test_xlsx_formula_missing_cache_hidden_sheet(self):
        from openpyxl import Workbook
        wb = Workbook()
        wb.active['A1'] = '=1+2'
        hidden = wb.create_sheet('隐藏')
        hidden.sheet_state = 'hidden'
        hidden['B2'] = '模拟隐藏证据'
        wb.save(self.root / 'sample.xlsx')
        self.sync()
        match = d.search(self.body(), 'formula=')['results'][0]
        self.assertIn('saved_value=MISSING', match['text'])
        self.assertIn('freshness=unknown', match['text'])
        self.assertEqual(d.search(self.body(), '隐藏证据')['results'][0]['location'], '隐藏!B2 (hidden)')

    def test_pdf_physical_page_and_scan_detection(self):
        from reportlab.pdfgen import canvas
        c = canvas.Canvas(str(self.root / 'text.pdf'))
        c.drawString(50, 700, 'SIMULATED first page')
        c.showPage()
        c.drawString(50, 700, 'payment 30 days')
        c.save()
        from pypdf import PdfWriter
        writer = PdfWriter()
        writer.add_blank_page(width=500, height=700)
        with (self.root / 'blank.pdf').open('wb') as f: writer.write(f)
        self.sync()
        self.assertEqual(d.search(self.body(), 'payment')['results'][0]['location'], 'PDF physical page 2')
        self.assertEqual(next(f for f in self.body()['files'] if f['path'] == 'blank.pdf')['status'], 'partial')

    def test_archive_symlink_destination(self):
        self.sync()
        real = self.base / 'real'
        real.mkdir()
        link = self.base / 'linked'
        try: link.symlink_to(real, target_is_directory=True)
        except OSError: self.skipTest('symlink not available')
        with self.assertRaises(ValueError): d.archive(self.body(), link, True)
        self.assertEqual(list(real.iterdir()), [])

    def test_saved_formula_cache(self):
        from openpyxl import Workbook
        import zipfile
        wb = Workbook()
        wb.active['A1'] = '=1+2'
        original = self.base / 'source.xlsx'
        wb.save(original)
        target = self.root / 'cache.xlsx'
        with zipfile.ZipFile(original) as src, zipfile.ZipFile(target, 'w') as out:
            for info in src.infolist():
                value = src.read(info.filename)
                if info.filename == 'xl/worksheets/sheet1.xml':
                    value = value.replace(b'<v></v>', b'<v>3</v>')
                out.writestr(info, value)
        self.sync()
        result = d.search(self.body(), 'formula=')['results'][0]['text']
        self.assertIn('saved_value=3', result)
        self.assertIn('freshness=unknown', result)

    def test_missing_ocr_does_not_claim_success(self):
        from pypdf import PdfWriter
        writer = PdfWriter()
        writer.add_blank_page(width=100, height=100)
        path = self.root / 'scan.pdf'
        with path.open('wb') as f: writer.write(f)
        with patch.object(d.subprocess, 'run', side_effect=FileNotFoundError('simulated missing OCR tool')):
            with self.assertRaises(FileNotFoundError): d.parse(path, ocr=True)

    def test_skill_packages_self_contained(self):
        import zipfile
        tool = SCRIPT.parents[3] / 'tools/build_skill.py'
        for target in ('standard', 'workbuddy'):
            proc = subprocess.run([sys.executable, str(tool), '--target', target, '--out', str(self.base / 'dist')], capture_output=True, text=True)
            self.assertEqual(proc.returncode, 0, proc.stderr)
            with zipfile.ZipFile(proc.stdout.strip()) as z:
                names = z.namelist()
                self.assertIn('organize-search-documents/scripts/docflow.py', names)
                self.assertIn('organize-search-documents/requirements.txt', names)
                skill = z.read('organize-search-documents/SKILL.md').decode()
                self.assertEqual('display_name:' in skill, target == 'workbuddy')
                z.extractall(self.base / target)
            script = self.base / target / 'organize-search-documents/scripts/docflow.py'
            run = subprocess.run([sys.executable, str(script), 'doctor'], capture_output=True, text=True)
            self.assertEqual(run.returncode, 0)
            self.assertEqual(json.loads(run.stdout)['version'], d.VERSION)

    def test_limits_missing_index_and_empty_query(self):
        with self.assertRaises(ValueError): d.snapshot(self.db)
        with patch.object(d, 'MAX_BYTES', 1):
            self.assertEqual(self.sync()['counts']['error'], 1)
        with self.assertRaises(ValueError): d.search(self.body(), ' ')
        with patch.object(d, 'MAX_CHARS', 2):
            with self.assertRaises(ValueError): d.parse(self.file)


if __name__ == '__main__':
    unittest.main(verbosity=2)
